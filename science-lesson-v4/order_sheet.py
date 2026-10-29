"""Science order sheet generator, based on the school's original card layout.

build(week_label, wc_text, lessons, out_path, style=1)
  lessons: {(day 0-4, period 1-5): dict(cls, room, topic, equip, size, sets)}
  Missing key = free period (empty card, as on the school sheet).
  style 1-5 = layout variant.
"""
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

TEACHER = "Patryk Tomczynski"
F = "Georgia"
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri"]
DAYS_LONG = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Per-class colours: (light tint, stronger accent). Both stay readable in greyscale.
CLASS_COL = {
    "7MLA": ("FCE4D6", "F4B183"),
    "8XP":  ("E2EFDA", "A9D18E"),
    "8AG":  ("DDEBF7", "9DC3E6"),
    "9AS":  ("FFF2CC", "FFD966"),
    "9XS":  ("EDE1F5", "C9A8E0"),
    "10S2": ("D5F0EC", "8FD3CB"),
    "10C2": ("FBE0EC", "F09CC0"),
    "11S2": ("F6E6D2", "DDB27C"),
    "12C":  ("DCE6F0", "8EA9C9"),
    "13D":  ("E6F2D2", "B5D98A"),
}
GREY = ("F2F2F2", "BFBFBF")

MED = Side(style="medium", color="000000")
NONE = Side(style=None)

CARD_W = [10.5, 5, 11.5, 8]


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def lines(text, chars):
    if not text:
        return 1
    return sum(max(1, math.ceil(len(t) / chars)) for t in text.split("\n"))


def build(week_label, wc_text, lessons, out_path, style=1):
    tab = style == 3                     # coloured tab column on each card
    gap = style in (4, 5)                # cutting gaps between cards
    ncol = 4 + (1 if tab else 0)
    wb = Workbook()
    ws = wb.active
    ws.title = "Order sheet"
    ws.sheet_view.showGridLines = False

    # Columns
    ws.column_dimensions["A"].width = 5
    starts = []
    col = 2
    for p in range(5):
        starts.append(col)
        widths = ([1.6] if tab else []) + CARD_W
        for w in widths:
            ws.column_dimensions[get_column_letter(col)].width = w
            col += 1
        if gap and p < 4:
            ws.column_dimensions[get_column_letter(col)].width = 1.2
            col += 1
    last_col = col - 1
    card_chars = 30                      # chars per line across a card, 11pt Georgia

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=starts[2] - 1)
    c = ws.cell(1, 1, f"Teacher:  {TEACHER}")
    c.font = Font(name=F, sz=13, b=True)
    c.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=1, start_column=starts[3], end_row=1, end_column=last_col)
    c = ws.cell(1, starts[3], f"{week_label}   {wc_text}")
    c.font = Font(name=F, sz=13, b=True)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 24

    # Period row
    c = ws.cell(2, 1, "Day")
    c.font = Font(name=F, sz=10, b=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for p in range(5):
        s = starts[p]
        ws.merge_cells(start_row=2, start_column=s, end_row=2, end_column=s + ncol - 1)
        c = ws.cell(2, s, f"Period {p + 1}")
        c.font = Font(name=F, sz=12, b=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(ncol):
            ws.cell(2, s + k).border = Border(bottom=MED)
    ws.row_dimensions[2].height = 20

    row = 3
    for d in range(5):
        r_h, r_rc, r_t, r_e, r_s = range(row, row + 5)

        t_l = e_l = 1
        for p in range(1, 6):
            L = lessons.get((d, p))
            if L:
                t_l = max(t_l, lines(L["topic"], card_chars))
                e_l = max(e_l, lines(L["equip"], card_chars - 4))
        ws.row_dimensions[r_h].height = 18
        ws.row_dimensions[r_rc].height = 19
        ws.row_dimensions[r_t].height = max(20, t_l * 15 + 6)
        ws.row_dimensions[r_e].height = max(62, e_l * 15 + 24)
        ws.row_dimensions[r_s].height = 19

        ws.merge_cells(start_row=r_h, start_column=1, end_row=r_s, end_column=1)
        c = ws.cell(r_h, 1, DAYS[d])
        c.font = Font(name=F, sz=12, b=True)
        c.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        for r in range(r_h, r_s + 1):
            ws.cell(r, 1).border = Border(left=MED, right=MED, top=MED if r == r_h else NONE,
                                          bottom=MED if r == r_s else NONE)

        for p in range(1, 6):
            L = lessons.get((d, p))
            free = L is None
            L = L or {}
            tint, accent = CLASS_COL.get(L.get("cls"), GREY)
            s0 = starts[p - 1]
            s = s0 + (1 if tab else 0)          # first content column

            lab = Font(name=F, sz=9.5, color="808080" if free else "000000")
            val = Font(name=F, sz=11, b=True)
            ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

            def put(r, c_, v, span=1, font=None, fl=None, al=None):
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=c_, end_row=r, end_column=c_ + span - 1)
                cell = ws.cell(r, c_, v)
                cell.font = font or lab
                cell.alignment = al or ctr
                if fl:
                    for k in range(span):
                        ws.cell(r, c_ + k).fill = fl
                return cell

            head = f"{DAYS[d]} P{p}"
            # Header row
            if style == 4 and not free:
                put(r_h, s, head, 4, Font(name=F, sz=11, b=True), fill(accent))
            else:
                put(r_h, s, head, 4, Font(name=F, sz=11, b=True, color="808080" if free else "000000"))

            # Room / Class
            if style == 2 and not free:
                for k in range(4):
                    ws.cell(r_rc, s + k).fill = fill(tint)
            put(r_rc, s, "Room:")
            put(r_rc, s + 1, L.get("room", ""), font=val)
            put(r_rc, s + 2, "Class:")
            put(r_rc, s + 3, L.get("cls", ""), font=val,
                fl=fill(tint) if style in (1, 3, 5) and not free else None)

            # Topic
            put(r_t, s, L.get("topic", ""), 4, Font(name=F, sz=11))

            # Equipment
            put(r_e, s, L.get("equip", ""), 4,
                Font(name=F, sz=11, b=True, i=L.get("equip") == "Nothing needed"),
                fill(tint) if style == 5 and not free else None,
                Alignment(horizontal="center", vertical="center", wrap_text=True))

            # Class size / sets
            put(r_s, s, "Class size:")
            put(r_s, s + 1, L.get("size", ""), font=val)
            put(r_s, s + 2, "No. of sets:")
            put(r_s, s + 3, L.get("sets", ""), font=val)

            # Coloured tab (style 3)
            if tab:
                ws.merge_cells(start_row=r_h, start_column=s0, end_row=r_s, end_column=s0)
                if not free:
                    ws.cell(r_h, s0).fill = fill(accent)
                    for r in range(r_h, r_s + 1):
                        ws.cell(r, s0).fill = fill(accent)

            # Outline only, no inner lines
            for r in range(r_h, r_s + 1):
                for k in range(ncol):
                    cc = s0 + k
                    ws.cell(r, cc).border = Border(
                        left=MED if k == 0 else NONE,
                        right=MED if k == ncol - 1 else NONE,
                        top=MED if r == r_h else NONE,
                        bottom=MED if r == r_s else NONE,
                    )

        row += 5
        if gap:
            ws.row_dimensions[row].height = 7
            row += 1

    # Colour key
    row += 0 if gap else 0
    ws.row_dimensions[row].height = 6
    row += 1
    used = sorted({L["cls"] for L in lessons.values()}, key=lambda x: (len(x), x))
    c = ws.cell(row, 1, "Key")
    c.font = Font(name=F, sz=9, b=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    slots = [st + (1 if tab else 0) + o for st in starts for o in (0, 2)]
    for cl, k_col in zip(used, slots):
        tint, accent = CLASS_COL.get(cl, GREY)
        c = ws.cell(row, k_col, cl)
        c.font = Font(name=F, sz=9, b=True)
        c.fill = fill(accent if style in (3, 4) else tint)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 16

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3, header=0.1, footer=0.1)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(last_col)}{row}"
    wb.save(out_path)
